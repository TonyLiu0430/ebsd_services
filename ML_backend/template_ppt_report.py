from __future__ import annotations

import io
import json
import math
import re
import statistics
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple, Union

from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ASSET_DIR = Path(__file__).resolve().parent / "assets"
MANIFEST_PATH = ASSET_DIR / "ebsd_report_template.json"
FONT_FACE = "Arial"
INK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 89, 89)
BLUE = "#2563eb"
RED = "#dc2626"
ORIENTATION_COLORS = ("#409eff", "#67c23a", "#e6a23c")
ORIENTATION_SAMPLE = "#409eff"
ORIENTATION_GOLDEN = "#e6a23c"
OVERRIDDEN_TEXT = RGBColor(0, 0, 255)
POSITION_RE = re.compile(r"^([A-Z]+)-([1-9]\d*)$")

Snapshot = Dict[str, Dict[str, Any]]
ReportData = Dict[str, Snapshot]


def _excel_number(column: str) -> int:
    result = 0
    for char in column:
        result = result * 26 + ord(char) - 64
    return result


def _excel_column(number: int) -> str:
    result = ""
    while number > 0:
        number -= 1
        result = chr(65 + number % 26) + result
        number //= 26
    return result


def _position_key(value: str) -> Tuple[int, int, str]:
    match = POSITION_RE.fullmatch(str(value).strip().upper())
    if not match:
        return (10**9, 10**9, str(value))
    return (int(match.group(2)), _excel_number(match.group(1)), str(value))


def _scan_position_key(value: str) -> Tuple[int, int, str]:
    match = POSITION_RE.fullmatch(str(value).strip().upper())
    if not match:
        return (10**9, 10**9, str(value))
    return (_excel_number(match.group(1)), int(match.group(2)), str(value))


def _positions(*snapshots: Snapshot) -> List[str]:
    return sorted({pos for snapshot in snapshots for pos in snapshot}, key=_position_key)


def _matrix_positions(positions: Iterable[str]) -> Tuple[List[str], List[int]]:
    parsed = [POSITION_RE.fullmatch(str(pos).strip().upper()) for pos in positions]
    parsed = [match for match in parsed if match]
    if not parsed:
        return [], []
    max_column = max(_excel_number(match.group(1)) for match in parsed)
    max_row = max(int(match.group(2)) for match in parsed)
    return [_excel_column(value) for value in range(1, max_column + 1)], list(range(1, max_row + 1))


def _finite(values: Any) -> List[float]:
    if not isinstance(values, list):
        return []
    return [float(value) for value in values if isinstance(value, (int, float)) and math.isfinite(float(value))]


def _orientation_labels(snapshot: Snapshot) -> List[str]:
    for item in snapshot.values():
        labels = item.get("orientation_labels")
        if isinstance(labels, list) and labels:
            return [str(value) for value in labels[:3]]
    return ["[001]", "[011]", "[111]"]


def _orientation_values(item: Optional[Dict[str, Any]], deviation: str, *, grains: bool = False) -> List[float]:
    prefix = "orientation_ratio_by_grains" if grains else "orientation_ratio"
    values = item.get(f"{prefix}({deviation})") if item else None
    if not isinstance(values, list):
        return [0.0, 0.0, 0.0]
    return [float(values[index] if index < len(values) else 0.0) * 100 for index in range(3)]


class TemplateDeck:
    def __init__(self) -> None:
        manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
        if manifest.get("schema_version") != 2:
            raise RuntimeError("unsupported EBSD PPT template manifest")
        self.prs = Presentation(ASSET_DIR / manifest["template_file"])
        self.manifest = manifest
        self.prototype_count = len(self.prs.slides)
        self.prototypes: Dict[str, Any] = {}
        for kind, spec in manifest["prototypes"].items():
            slide = self.prs.slides[int(spec["slide_number"]) - 1]
            if not any(shape.name == spec["prototype_shape"] for shape in slide.shapes):
                raise RuntimeError(f"invalid clean template prototype: {kind}")
            self.prototypes[kind] = slide

    def add_slide(self, title: str, subtitle: str = "", kind: str = "generic"):
        prototype = self.prototypes.get(kind, self.prototypes["generic"])
        slide = self.prs.slides.add_slide(prototype.slide_layout)
        for shape in prototype.shapes:
            slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
        title_shape = self.shape(slide, "EBSD_TITLE")
        _replace_shape_text(title_shape, title)
        heading = self.manifest["page_heading"]
        title_shape.left = Inches(heading["title_left_in"])
        title_shape.top = Inches(heading["title_top_in"])
        title_shape.width = Inches(heading["title_width_in"])
        title_shape.height = Inches(heading["title_height_in"])
        subtitle_shape = next((shape for shape in slide.shapes if shape.name == "EBSD_SUBTITLE"), None)
        if subtitle_shape is not None:
            _replace_shape_text(subtitle_shape, subtitle)
            subtitle_shape.left = Inches(heading["subtitle_left_in"])
            subtitle_shape.top = Inches(heading["subtitle_top_in"])
            subtitle_shape.width = Inches(heading["subtitle_width_in"])
            subtitle_shape.height = Inches(heading["subtitle_height_in"])
        return slide

    @staticmethod
    def shape(slide, name: str):
        shape = next((candidate for candidate in slide.shapes if candidate.name == name), None)
        if shape is None:
            raise RuntimeError(f"template shape is missing: {name}")
        return shape

    def remove_shape(self, slide, name: str) -> None:
        shape = next((candidate for candidate in slide.shapes if candidate.name == name), None)
        if shape is not None:
            shape._element.getparent().remove(shape._element)

    def apply_table_template(self, table, kind: str, name: str) -> None:
        prototype = self.shape(self.prototypes[kind], name).table
        current = table._tbl.tblPr
        current.getparent().replace(current, deepcopy(prototype._tbl.tblPr))

    @property
    def content(self) -> Tuple[int, int, int, int]:
        box = self.manifest["generic_content_box"]
        return tuple(Inches(box[key]) for key in ("left_in", "top_in", "width_in", "height_in"))  # type: ignore[return-value]

    def save(self) -> bytes:
        for _ in range(self.prototype_count):
            slide_id = self.prs.slides._sldIdLst[0]
            self.prs.part.drop_rel(slide_id.rId)
            del self.prs.slides._sldIdLst[0]
        output = io.BytesIO()
        self.prs.save(output)
        return output.getvalue()


def _replace_shape_text(shape, text: str) -> None:
    if not shape.has_text_frame:
        return
    runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
    if runs:
        runs[0].text = str(text)
        for run in runs[1:]:
            run.text = ""
        return
    shape.text_frame.text = str(text)


def _set_scan_title(deck: TemplateDeck, slide, sample: str, position: str, reanalysis: bool) -> None:
    shape = deck.shape(slide, "EBSD_TITLE")
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    base = paragraph.add_run()
    base.text = f"{sample} {position}"
    base.font.name = FONT_FACE
    base.font.size = Pt(24)
    base.font.bold = True
    base.font.color.rgb = INK
    if reanalysis:
        suffix = paragraph.add_run()
        suffix.text = " (Reanalysis)"
        suffix.font.name = FONT_FACE
        suffix.font.size = Pt(24)
        suffix.font.bold = True
        suffix.font.color.rgb = RGBColor(0, 0, 255)


def _add_text(slide, text: str, left: int, top: int, width: int, height: int, *, size: int = 10, bold: bool = False, color: RGBColor = INK, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = FONT_FACE
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def _set_cell(cell, text: str, *, header: bool = False, size: int = 8, color: Optional[RGBColor] = None) -> None:
    cell.margin_left = Inches(0.04)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.text = ""
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = FONT_FACE
    run.font.size = Pt(size)
    run.font.bold = header
    run.font.color.rgb = color or INK


def _add_native_matrix(
    deck: TemplateDeck,
    slide,
    positions: Sequence[str],
    value_for_position,
    *,
    header_column_width: float = 0.992,
    data_column_width: float = 0.992,
    row_height: float = 0.406,
    style_kind: str = "orientation_pixels",
    style_shape: str = "EBSD_TABLE_20",
    font_size: int = 18,
) -> None:
    columns, rows = _matrix_positions(positions)
    if not columns or not rows:
        left, top, width, height = deck.content
        _add_text(slide, "沒有可輸出的動態位置資料", left, top, width, height, size=16, color=MUTED, align=PP_ALIGN.CENTER)
        return
    left, top, _, _ = deck.content
    table_width = Inches(header_column_width + len(columns) * data_column_width)
    table_height = Inches((len(rows) + 1) * row_height)
    table = slide.shapes.add_table(len(rows) + 1, len(columns) + 1, left, top, table_width, table_height).table
    deck.apply_table_template(table, style_kind, style_shape)
    table.columns[0].width = Inches(header_column_width)
    for index in range(1, len(columns) + 1):
        table.columns[index].width = Inches(data_column_width)
    for row in table.rows:
        row.height = Inches(row_height)
    _set_cell(table.cell(0, 0), "位置", header=True, size=font_size)
    for column_index, column in enumerate(columns, start=1):
        _set_cell(table.cell(0, column_index), column, header=True, size=font_size)
    position_set = set(positions)
    for row_index, row in enumerate(rows, start=1):
        _set_cell(table.cell(row_index, 0), str(row), header=True, size=font_size)
        for column_index, column in enumerate(columns, start=1):
            position = f"{column}-{row}"
            _set_cell(table.cell(row_index, column_index), value_for_position(position) if position in position_set else "", size=font_size)


def _add_picture_contain(slide, data: Union[bytes, io.BytesIO], left: int, top: int, width: int, height: int) -> None:
    stream = data if isinstance(data, io.BytesIO) else io.BytesIO(data)
    stream.seek(0)
    with Image.open(stream) as image:
        image_width, image_height = image.size
    stream.seek(0)
    ratio = min(width / max(image_width, 1), height / max(image_height, 1))
    draw_width = int(image_width * ratio)
    draw_height = int(image_height * ratio)
    slide.shapes.add_picture(stream, left + (width - draw_width) // 2, top + (height - draw_height) // 2, draw_width, draw_height)


def _add_captioned_picture(slide, caption: str, image: Optional[bytes], left: float, top: float, width: float, height: float) -> None:
    _add_text(
        slide,
        caption,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(0.28),
        size=12,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    if image:
        _add_picture_contain(
            slide,
            image,
            Inches(left),
            Inches(top + 0.30),
            Inches(width),
            Inches(height - 0.30),
        )


def _add_grain_summary_table(slide, item: Dict[str, Any], left: float, top: float, width: float, height: float) -> None:
    grains = _finite(item.get("grains"))
    values = (
        ("Average", f"{statistics.mean(grains):.3f}" if grains else "—"),
        ("Max", f"{max(grains):.3f}" if grains else "—"),
        ("Quantity", str(len(grains))),
    )
    table = slide.shapes.add_table(4, 2, Inches(left), Inches(top), Inches(width), Inches(height)).table
    table.columns[0].width = Inches(width * 0.57)
    table.columns[1].width = Inches(width * 0.43)
    _set_cell(table.cell(0, 0), "Grain size", header=True, size=9)
    _set_cell(table.cell(0, 1), "(um)", header=True, size=9)
    for row_index, (label, value) in enumerate(values, start=1):
        _set_cell(table.cell(row_index, 0), label, size=8)
        _set_cell(table.cell(row_index, 1), value, size=8)


def _add_ratio_table(slide, item: Dict[str, Any], deviation: str, left: float, top: float, width: float, height: float) -> None:
    labels = _orientation_labels({"_": item})
    values = _orientation_values(item, deviation)
    table = slide.shapes.add_table(2, 3, Inches(left), Inches(top), Inches(width), Inches(height)).table
    for column in table.columns:
        column.width = Inches(width / 3)
    for column_index in range(3):
        _set_cell(table.cell(0, column_index), f"<{labels[column_index].strip('[]')}>", header=True, size=9)
        table.cell(0, column_index).fill.solid()
        table.cell(0, column_index).fill.fore_color.rgb = RGBColor(70, 167, 193)
        _set_cell(table.cell(1, column_index), f"{values[column_index]:.1f}%", size=9)


def _add_scan_point_pages(
    deck: TemplateDeck,
    sample: str,
    snapshot: Snapshot,
    images_by_position: Dict[str, Dict[str, Any]],
    overridden_positions: Sequence[str],
    ipf_legend: Optional[bytes],
    histogram_settings: Dict[str, Any],
) -> None:
    overridden = set(overridden_positions)
    for position in sorted(snapshot, key=_scan_position_key):
        images = images_by_position.get(position)
        if not images:
            continue
        item = snapshot.get(position, {})

        overview = deck.add_slide("", kind="generic")
        _set_scan_title(deck, overview, sample, position, position in overridden)
        _add_captioned_picture(overview, "Band contrast", images.get("band_contrast"), 0.45, 0.65, 3.55, 2.45)
        _add_captioned_picture(overview, "Grain size (D)", images.get("grain_map"), 4.10, 0.65, 3.55, 2.45)
        grains = _finite(item.get("grains"))
        histogram = _grain_chart(
            grains,
            [],
            "ratio",
            _histogram_edges(grains, histogram_settings),
        )
        _add_captioned_picture(overview, "Grain-size distribution", histogram.getvalue(), 7.78, 0.65, 3.35, 2.45)
        _add_grain_summary_table(overview, item, 11.15, 1.02, 1.72, 1.55)

        ipf_maps = images.get("ipf_maps", {})
        for index, direction in enumerate(("X", "Y", "Z")):
            _add_captioned_picture(
                overview,
                f"IPF-{direction}",
                ipf_maps.get(direction),
                0.45 + index * 4.05,
                3.35,
                3.75,
                3.35,
            )
        if ipf_legend:
            _add_captioned_picture(overview, "IPF color key", ipf_legend, 12.10, 5.22, 0.95, 1.20)

        texture = deck.add_slide("", kind="generic")
        _set_scan_title(deck, texture, sample, position, position in overridden)
        _add_text(texture, "Preferred orientation ratio", Inches(0.35), Inches(0.58), Inches(3.0), Inches(0.28), size=13, bold=True)
        _add_text(texture, "Direction: Z", Inches(0.45), Inches(0.88), Inches(1.5), Inches(0.24), size=10)
        _add_captioned_picture(texture, "Range: 20°", images.get("preferred_20"), 1.95, 0.72, 4.15, 2.55)
        _add_captioned_picture(texture, "Range: 15°", images.get("preferred_15"), 6.65, 0.72, 4.15, 2.55)
        _add_ratio_table(texture, item, "20%", 2.25, 3.18, 3.55, 0.72)
        _add_ratio_table(texture, item, "15%", 6.95, 3.18, 3.55, 0.72)
        if ipf_legend:
            _add_captioned_picture(texture, "Orientation key", ipf_legend, 11.25, 1.15, 1.45, 1.75)
        _add_captioned_picture(texture, "Pole Figure", images.get("pole_figure"), 0.25, 4.05, 6.25, 3.15)
        _add_captioned_picture(texture, "Inverse Pole Figure", images.get("inverse_pole_figure"), 6.75, 4.05, 6.25, 3.15)


def _prepare_ipf_matrix(deck: TemplateDeck, slide, columns: Sequence[str], rows: Sequence[int]):
    """Replace the fixed 3x3 template grid with one unscaled, dynamic matrix.

    The template's data-cell size is intentionally retained.  Matrices larger
    than 3x3 therefore extend beyond the slide instead of being split across
    several slides, so all positions remain together for comparison.
    """
    layout = deck.manifest["ipf"]
    grid_shape = deck.shape(slide, "EBSD_IPF_GRID")
    prototype_table = grid_shape.table
    data_columns = [column for index, column in enumerate(prototype_table.columns) if index > 0]
    data_rows = [row for index, row in enumerate(prototype_table.rows) if index > 0]
    header_column_width = prototype_table.columns[0].width
    data_column_width = round(
        sum(column.width for column in data_columns)
        / max(1, len(data_columns))
    )
    header_row_height = prototype_table.rows[0].height
    data_row_height = round(
        sum(row.height for row in data_rows)
        / max(1, len(data_rows))
    )
    left = Inches(layout["grid_left_in"])
    top = Inches(layout["grid_top_in"])
    deck.remove_shape(slide, "EBSD_IPF_GRID")

    table = slide.shapes.add_table(
        len(rows) + 1,
        len(columns) + 1,
        left,
        top,
        header_column_width + len(columns) * data_column_width,
        header_row_height + len(rows) * data_row_height,
    ).table
    deck.apply_table_template(table, "ipf", "EBSD_IPF_GRID")
    table.columns[0].width = header_column_width
    for column_index, column in enumerate(table.columns):
        if column_index > 0:
            column.width = data_column_width
    table.rows[0].height = header_row_height
    for row_index, row in enumerate(table.rows):
        if row_index > 0:
            row.height = data_row_height

    _set_cell(table.cell(0, 0), "", size=18)
    for column_index, column in enumerate(columns, start=1):
        _set_cell(table.cell(0, column_index), column, size=18)
    for row_index, row in enumerate(rows, start=1):
        _set_cell(table.cell(row_index, 0), str(row), header=True, size=18)
        for column_index in range(1, len(columns) + 1):
            _set_cell(table.cell(row_index, column_index), "", size=18)

    horizontal_padding = Inches(0.07)
    vertical_padding = Inches(0.02)
    return (
        left + header_column_width,
        top + header_row_height,
        data_column_width,
        data_row_height,
        horizontal_padding,
        vertical_padding,
    )


def _add_ipf_pages(deck: TemplateDeck, selected_sample: str, version_label: str, directions: Dict[str, Dict[str, bytes]], legend: Optional[bytes]) -> None:
    for direction, images in directions.items():
        columns, rows = _matrix_positions(images)
        if not columns or not rows:
            continue
        slide = deck.add_slide(
            f"{selected_sample}-Orientation maps with IPF {direction.upper()}",
            kind="ipf",
        )
        image_left, image_top, cell_width, cell_height, pad_x, pad_y = _prepare_ipf_matrix(
            deck, slide, columns, rows
        )
        for row_index, row in enumerate(rows):
            for column_index, column in enumerate(columns):
                image = images.get(f"{column}-{row}")
                if not image:
                    continue
                _add_picture_contain(
                    slide,
                    image,
                    image_left + column_index * cell_width + pad_x,
                    image_top + row_index * cell_height + pad_y,
                    cell_width - 2 * pad_x,
                    cell_height - 2 * pad_y,
                )


def _histogram_edges(all_values: Sequence[float], settings: Dict[str, Any]) -> List[float]:
    maximum = 200.0
    minimum = 0.0 if settings.get("start_mode") == "zero" else (math.floor(min(all_values) * 10) / 10 if all_values else 7.1)
    span = max(maximum - minimum, 0.1)
    if settings.get("width_mode") == "absolute":
        width = max(float(settings.get("width_absolute", 10.0)), 0.1)
    else:
        width = max(span * float(settings.get("width_ratio", 10.0)) / 100.0, 0.1)
    count = max(1, min(80, math.ceil(span / width)))
    return [minimum + (maximum - minimum) * index / count for index in range(count + 1)]


def _hist(values: Sequence[float], edges: Sequence[float], *, ratio: bool, area: bool) -> List[float]:
    result = [0.0] * (len(edges) - 1)
    total = 0.0
    for value in values:
        if not math.isfinite(value):
            continue
        index = min(len(result) - 1, max(0, int((value - edges[0]) / max(edges[-1] - edges[0], 1e-9) * len(result))))
        weight = math.pi * (value / 2.0) ** 2 if area else 1.0
        result[index] += weight
        total += weight
    if (ratio or area) and total > 0:
        result = [value / total * 100 for value in result]
    return result


def _grain_chart(sample: Sequence[float], golden: Sequence[float], mode: str, edges: Sequence[float]) -> io.BytesIO:
    figure = Figure(figsize=(3.75, 2.5), dpi=220)
    canvas = FigureCanvasAgg(figure)
    axis = figure.add_subplot(111)
    if mode == "cdf":
        for values, color, linestyle, label in ((golden, BLUE, "-", "Golden"), (sample, RED, "--", "Sample")):
            ordered = sorted(values)
            if ordered:
                axis.plot(ordered, [(index + 1) / len(ordered) * 100 for index in range(len(ordered))], color=color, linewidth=1.8, linestyle=linestyle, label=label)
        axis.set_ylim(0, 100)
        axis.set_ylabel("Cumulative %")
    else:
        ratio = mode == "ratio"
        area = mode == "area"
        sample_series = _hist(sample, edges, ratio=ratio, area=area)
        golden_series = _hist(golden, edges, ratio=ratio, area=area)
        width = edges[1] - edges[0]
        axis.bar(edges[:-1], golden_series, width=width, align="edge", color=BLUE, alpha=0.22, edgecolor=BLUE, label="Golden")
        axis.bar(edges[:-1], sample_series, width=width, align="edge", color=RED, alpha=0.2, edgecolor=RED, linestyle="--", label="Sample")
        if ratio:
            # Requirement: histogram ratio y-axis always tops out at 100%.
            axis.set_ylim(0, 100)
            axis.set_ylabel("Ratio %")
        elif area:
            axis.set_ylim(0, 100)
            axis.set_ylabel("Area %")
        else:
            axis.set_ylabel("Count")
    axis.set_xlim(0, 200)
    axis.set_xlabel("Grain (μm)")
    axis.grid(axis="y", color="#eeeeee", linewidth=0.7)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.tick_params(labelsize=7)
    axis.legend(loc="upper right", fontsize=5.5, frameon=False)
    axis.xaxis.label.set_size(8)
    axis.yaxis.label.set_size(8)
    figure.tight_layout(pad=0.35)
    output = io.BytesIO()
    canvas.print_png(output)
    output.seek(0)
    return output


def _add_grain_chart_pages(deck: TemplateDeck, selected: Snapshot, golden: Snapshot, selected_label: str, golden_label: str, settings: Dict[str, Any]) -> None:
    positions = _positions(selected, golden)
    all_values = [value for snapshot in (selected, golden) for item in snapshot.values() for value in _finite(item.get("grains"))]
    edges = _histogram_edges(all_values, settings)
    modes = (
        ("cdf", "Grain cumulative curve"),
        ("count", "Grain histogram (Count)"),
        ("ratio", "Grain histogram (Ratio, Y max 100%)"),
        ("area", "Area-weighted grain histogram"),
    )
    matrix_columns, matrix_rows = _matrix_positions(positions)
    position_set = set(positions)
    if not matrix_columns or not matrix_rows:
        return
    for mode, title in modes:
        slide = deck.add_slide(title, kind="ipf")
        image_left, image_top, cell_width, cell_height, pad_x, pad_y = _prepare_ipf_matrix(
            deck, slide, matrix_columns, matrix_rows
        )
        for row_index, row in enumerate(matrix_rows):
            for column_index, column in enumerate(matrix_columns):
                position = f"{column}-{row}"
                if position not in position_set:
                    continue
                png = _grain_chart(
                    _finite(selected.get(position, {}).get("grains")),
                    _finite(golden.get(position, {}).get("grains")),
                    mode,
                    edges,
                )
                _add_picture_contain(
                    slide,
                    png,
                    image_left + column_index * cell_width + pad_x,
                    image_top + row_index * cell_height + pad_y,
                    cell_width - 2 * pad_x,
                    cell_height - 2 * pad_y,
                )


def _complete_cell(item: Optional[Dict[str, Any]]) -> str:
    grains = _finite(item.get("grains")) if item else []
    if not grains:
        return "—"
    labels = _orientation_labels({"_": item or {}})
    ratio20 = _orientation_values(item, "20%")
    ratio15 = _orientation_values(item, "15%")
    return "\n".join((
        f"Avg {statistics.mean(grains):.2f}",
        f"Max {max(grains):.2f}",
        f"Qty {len(grains)}",
        "20° " + " ".join(f"{labels[i]} {ratio20[i]:.1f}%" for i in range(3)),
        "15° " + " ".join(f"{labels[i]} {ratio15[i]:.1f}%" for i in range(3)),
    ))


def _add_complete_table(
    deck: TemplateDeck,
    sample: str,
    version: str,
    snapshot: Snapshot,
    overridden_positions: Sequence[str] = (),
) -> None:
    slide = deck.add_slide("Complete EBSD data", f"{sample} · {version}")
    positions = sorted(snapshot, key=_position_key)
    overridden_set = set(overridden_positions)
    labels = _orientation_labels(snapshot)
    columns = 10
    column_width = deck.manifest["orientation_table"]["column_width_in"]
    row_height = deck.manifest["orientation_table"]["row_height_in"]
    left = Inches(deck.manifest["orientation_table"]["left_in"])
    top = Inches(deck.manifest["orientation_table"]["top_20_in"])
    table = slide.shapes.add_table(
        len(positions) + 2,
        columns,
        left,
        top,
        Inches(column_width * columns),
        Inches(row_height * (len(positions) + 2)),
    ).table
    deck.apply_table_template(table, "orientation_pixels", "EBSD_TABLE_20")
    for column in table.columns:
        column.width = Inches(column_width)
    for row in table.rows:
        row.height = Inches(row_height)
    for start, end, text in ((1, 3, "Grain Size"), (4, 6, "20°"), (7, 9, "15°")):
        table.cell(0, start).merge(table.cell(0, end))
        _set_cell(table.cell(0, start), text, header=True, size=18)
    _set_cell(table.cell(0, 0), "EBSD", header=True, size=18)
    headers = ["Position", "Average", "Max", "Quantity", *labels, *labels]
    for column_index, header in enumerate(headers):
        _set_cell(table.cell(1, column_index), header, header=True, size=12)
    for row_index, position in enumerate(positions, start=2):
        item = snapshot.get(position, {})
        grains = _finite(item.get("grains"))
        ratio20 = _orientation_values(item, "20%")
        ratio15 = _orientation_values(item, "15%")
        values = [
            position,
            f"{statistics.mean(grains):.2f}" if grains else "—",
            f"{max(grains):.2f}" if grains else "—",
            str(len(grains)),
            *(f"{value:.1f}%" for value in ratio20),
            *(f"{value:.1f}%" for value in ratio15),
        ]
        for column_index, value in enumerate(values):
            _set_cell(
                table.cell(row_index, column_index),
                value,
                header=column_index == 0,
                size=14,
                color=OVERRIDDEN_TEXT if column_index > 0 and position in overridden_set else None,
            )


def _add_orientation_tables(
    deck: TemplateDeck,
    snapshot: Snapshot,
    *,
    grains: bool,
    overridden_positions: Sequence[str] = (),
) -> None:
    labels = _orientation_labels(snapshot)
    overridden_set = set(overridden_positions)
    kind = "orientation_grains" if grains else "orientation_pixels"
    title = "Orientation ratio by grains" if grains else "Orientation ratio by pixels"
    slide = deck.add_slide(title, kind=kind)
    deck.remove_shape(slide, "EBSD_TABLE_20")
    deck.remove_shape(slide, "EBSD_TABLE_15")
    columns, rows = _matrix_positions(snapshot)
    if not columns or not rows:
        return
    spec = deck.manifest["orientation_table"]
    row_height = float(spec["row_height_in"])
    table_rows = len(rows) + 2
    top20 = float(spec["top_20_in"])
    top15 = top20 + table_rows * row_height + float(spec["section_gap_in"])
    subtitle20 = deck.shape(slide, "EBSD_SUBTITLE_20")
    subtitle15 = deck.shape(slide, "EBSD_SUBTITLE_15")
    _replace_shape_text(subtitle20, "Prefer orientation ratio (Misorientation angle 20°)")
    _replace_shape_text(subtitle15, "Prefer orientation ratio (Misorientation angle 15°)")
    subtitle15.top = Inches(top15 - 0.477)

    for deviation, top in (("20%", top20), ("15%", top15)):
        column_count = 1 + len(columns) * 3
        table = slide.shapes.add_table(
            table_rows,
            column_count,
            Inches(spec["left_in"]),
            Inches(top),
            Inches(float(spec["column_width_in"]) * column_count),
            Inches(row_height * table_rows),
        ).table
        deck.apply_table_template(table, kind, "EBSD_TABLE_20" if deviation == "20%" else "EBSD_TABLE_15")
        for column in table.columns:
            column.width = Inches(spec["column_width_in"])
        for row in table.rows:
            row.height = Inches(row_height)
        _set_cell(table.cell(0, 0), "EBSD", header=True, size=18)
        for axis_index, label in enumerate(labels):
            start = 1 + axis_index * len(columns)
            end = start + len(columns) - 1
            if end > start:
                table.cell(0, start).merge(table.cell(0, end))
            _set_cell(table.cell(0, start), label.strip("[]"), header=True, size=18)
            for column_index, column in enumerate(columns):
                _set_cell(table.cell(1, start + column_index), column, header=True, size=18)
        for row_index, row in enumerate(rows, start=2):
            _set_cell(table.cell(row_index, 0), str(row), header=True, size=18)
            for axis_index in range(3):
                for column_index, column in enumerate(columns):
                    position = f"{column}-{row}"
                    values = _orientation_values(snapshot.get(position), deviation, grains=grains)
                    _set_cell(
                        table.cell(row_index, 1 + axis_index * len(columns) + column_index),
                        f"{values[axis_index]:.1f}%" if position in snapshot else "",
                        size=18,
                        color=OVERRIDDEN_TEXT if position in overridden_set else None,
                    )


def _add_statistics(
    deck: TemplateDeck,
    statistics_by_min_size: Dict[int, Snapshot],
    sample: str,
    overridden_positions: Sequence[str] = (),
) -> None:
    overridden_set = set(overridden_positions)
    metric_specs = (
        ("Average", "EBSD_STATS_AVERAGE", lambda values: f"{statistics.mean(values):.2f}" if values else "—"),
        ("Max", "EBSD_STATS_MAX", lambda values: f"{max(values):.2f}" if values else "—"),
        ("Quantity", "EBSD_STATS_QUANTITY", lambda values: str(len(values))),
    )
    for threshold in sorted(statistics_by_min_size):
        snapshot = statistics_by_min_size[threshold]
        columns, rows = _matrix_positions(snapshot)
        for metric_name, style_shape, formatter in metric_specs:
            slide = deck.add_slide(f"{sample}-  Grain Size (min grain size = {threshold} )", kind="statistics")
            for name in ("EBSD_STATS_AVERAGE", "EBSD_STATS_MAX", "EBSD_STATS_QUANTITY"):
                deck.remove_shape(slide, name)
            if not columns or not rows:
                continue
            spec = deck.manifest["statistics_table"]
            row_count = len(rows) + 2
            column_count = len(columns) + 1
            table = slide.shapes.add_table(
                row_count,
                column_count,
                Inches(spec["left_in"]),
                Inches(spec["top_in"]),
                Inches(float(spec["column_width_in"]) * column_count),
                Inches(float(spec["row_height_in"]) * row_count),
            ).table
            deck.apply_table_template(table, "statistics", style_shape)
            for column in table.columns:
                column.width = Inches(spec["column_width_in"])
            for row in table.rows:
                row.height = Inches(spec["row_height_in"])
            if column_count > 1:
                table.cell(0, 0).merge(table.cell(0, column_count - 1))
            _set_cell(table.cell(0, 0), f"Grain Size {metric_name}", header=True, size=18)
            _set_cell(table.cell(1, 0), "" if metric_name == "Quantity" else "(um)", header=True, size=18)
            for column_index, column in enumerate(columns, start=1):
                _set_cell(table.cell(1, column_index), column, header=True, size=18)
            for row_index, row in enumerate(rows, start=2):
                _set_cell(table.cell(row_index, 0), str(row), header=True, size=18)
                for column_index, column in enumerate(columns, start=1):
                    position = f"{column}-{row}"
                    grains = _finite(snapshot.get(position, {}).get("grains"))
                    _set_cell(
                        table.cell(row_index, column_index),
                        formatter(grains) if position in snapshot else "",
                        size=18,
                        color=OVERRIDDEN_TEXT if position in overridden_set else None,
                    )


def _triangle_chart(
    position: str,
    selected_item: Optional[Dict[str, Any]],
    golden_item: Optional[Dict[str, Any]],
    deviation: str,
    labels: List[str],
) -> io.BytesIO:
    """Render one scan position using the same triangle system as TriangleChart.vue."""
    figure = Figure(figsize=(3.75, 2.5), dpi=220)
    canvas = FigureCanvasAgg(figure)
    axis = figure.add_subplot(111)
    vertices = [(0.0, 1.0), (math.sqrt(3) / 2, -0.5), (-math.sqrt(3) / 2, -0.5)]

    def points(values: Sequence[float]) -> List[Tuple[float, float]]:
        return [
            (vertices[index][0] * float(value) / 50.0, vertices[index][1] * float(value) / 50.0)
            for index, value in enumerate(values[:3])
        ]

    for level in (10, 25, 50):
        polygon = points([level, level, level])
        polygon.append(polygon[0])
        axis.plot([point[0] for point in polygon], [point[1] for point in polygon], color="#d1d5db", linewidth=0.8)
        label_point = points([level, 0, 0])[0]
        axis.text(label_point[0] + 0.03, label_point[1], f"{level}%", fontsize=5.5, color="#9ca3af", va="center")
    for vertex, label in zip(vertices, labels):
        axis.plot([0, vertex[0]], [0, vertex[1]], color="#9ca3af", linewidth=0.8)
        axis.text(vertex[0] * 1.18, vertex[1] * 1.18, label, fontsize=8, fontweight="bold", color="#374151", ha="center", va="center")

    series = (
        ("Sample", selected_item, ORIENTATION_SAMPLE, False),
        ("Golden", golden_item, ORIENTATION_GOLDEN, True),
    )
    for source, item, color, dashed in series:
        if item is None:
            continue
        values = _orientation_values(item, deviation)
        polygon = points(values)
        closed = polygon + polygon[:1]
        axis.plot(
            [point[0] for point in closed],
            [point[1] for point in closed],
            color=color,
            linewidth=1.35,
            linestyle="--" if dashed else "-",
            label=source,
        )
        axis.fill([point[0] for point in closed], [point[1] for point in closed], color=color, alpha=0.06)
        axis.scatter([point[0] for point in polygon], [point[1] for point in polygon], color=color, s=8)

    axis.set_title(position, fontsize=9, fontweight="bold", pad=2)
    axis.set_aspect("equal", adjustable="box")
    axis.set_xlim(-1.42, 1.42)
    axis.set_ylim(-1.18, 1.34)
    axis.axis("off")
    handles, legend_labels = axis.get_legend_handles_labels()
    if handles:
        axis.legend(handles, legend_labels, loc="upper center", bbox_to_anchor=(0.5, -0.01), ncol=2, fontsize=5.5, frameon=False)
    figure.tight_layout(pad=0.18)
    output = io.BytesIO()
    canvas.print_png(output)
    output.seek(0)
    return output


def _add_triangle_pages(deck: TemplateDeck, selected: Snapshot, golden: Snapshot) -> None:
    positions = _positions(selected, golden)
    labels = _orientation_labels(selected)
    columns, rows = _matrix_positions(positions)
    if not columns or not rows:
        return
    position_set = set(positions)
    for deviation in ("20%", "15%"):
        slide = deck.add_slide(
            f"晶粒取向比例 - Misorientation {deviation}",
            "Sample 藍色實線 / Golden 橘色虛線" if golden else "Sample 藍色實線",
            kind="ipf",
        )
        image_left, image_top, cell_width, cell_height, pad_x, pad_y = _prepare_ipf_matrix(
            deck, slide, columns, rows
        )
        for row_index, row in enumerate(rows):
            for column_index, column in enumerate(columns):
                position = f"{column}-{row}"
                if position not in position_set:
                    continue
                png = _triangle_chart(
                    position,
                    selected.get(position),
                    golden.get(position),
                    deviation,
                    labels,
                )
                _add_picture_contain(
                    slide,
                    png,
                    image_left + column_index * cell_width + pad_x,
                    image_top + row_index * cell_height + pad_y,
                    cell_width - 2 * pad_x,
                    cell_height - 2 * pad_y,
                )


def _line_chart(columns: Sequence[str], selected_values: Sequence[float], golden_values: Optional[Sequence[float]], axis_label: str) -> io.BytesIO:
    figure = Figure(figsize=(4.0, 2.0), dpi=220)
    canvas = FigureCanvasAgg(figure)
    axis = figure.add_subplot(111)
    x = list(range(len(columns)))
    axis.plot(x, selected_values, color=ORIENTATION_SAMPLE, marker="o", markersize=3.5, linewidth=1.7, label="Sample")
    if golden_values is not None:
        axis.plot(x, golden_values, color=ORIENTATION_GOLDEN, marker="o", markersize=3.5, linewidth=1.7, linestyle="--", label="Golden")
    axis.set_xticks(x, columns, rotation=0)
    axis.set_ylim(0, 100)
    axis.set_ylabel("Ratio %", fontsize=6)
    axis.set_title(axis_label, fontsize=8, pad=3)
    axis.grid(axis="y", color="#eeeeee")
    axis.tick_params(labelsize=5.5)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    figure.tight_layout(pad=0.3)
    output = io.BytesIO()
    canvas.print_png(output)
    output.seek(0)
    return output


def _add_orientation_lines(deck: TemplateDeck, selected: Snapshot, golden: Snapshot) -> None:
    positions = _positions(selected, golden)
    labels = _orientation_labels(selected)
    columns, rows = _matrix_positions(positions)
    for deviation in ("20%", "15%"):
        slide = deck.add_slide(
            f"晶粒取向折線圖 - Misorientation {deviation}",
            "Sample 藍色實線 / Golden 橘色虛線" if golden else "Sample 藍色實線",
        )
        left, top, width, height = deck.content
        cell_width = width // 3
        cell_height = height // max(1, len(rows))
        for row_index, row in enumerate(rows):
            for axis_index, label in enumerate(labels):
                selected_values = [
                    _orientation_values(selected.get(f"{column}-{row}"), deviation)[axis_index]
                    for column in columns
                ]
                golden_values = (
                    [
                        _orientation_values(golden.get(f"{column}-{row}"), deviation)[axis_index]
                        for column in columns
                    ]
                    if golden
                    else None
                )
                png = _line_chart(columns, selected_values, golden_values, f"Row {row} · {label}")
                _add_picture_contain(
                    slide,
                    png,
                    left + axis_index * cell_width,
                    top + row_index * cell_height,
                    cell_width,
                    cell_height,
                )


def _analysis_color(value: float, thresholds: Sequence[float]) -> RGBColor:
    low, high = (list(thresholds) + [30.0, 70.0])[:2]
    absolute = abs(value)
    if absolute <= low:
        return RGBColor(16, 185, 129)
    if absolute >= high:
        return RGBColor(239, 68, 68)
    return RGBColor(245, 158, 11)


def _add_feature_analysis(deck: TemplateDeck, analysis: Dict[str, Dict[str, float]], selected_sample: str, features: Sequence[str], thresholds: Sequence[float]) -> None:
    values = analysis.get(selected_sample, {})
    if not values or not features:
        return
    slide = deck.add_slide("Feature analysis", "Features: " + ", ".join(features))
    positions = sorted(values, key=_position_key)
    _add_native_matrix(deck, slide, positions, lambda position: f"{values[position] * 100:.1f}%", data_column_width=0.9, row_height=0.42)
    # Recolour only data cells, preserving the native editable table.
    table_shape = next(shape for shape in reversed(slide.shapes) if getattr(shape, "has_table", False))
    table = table_shape.table
    columns, rows = _matrix_positions(positions)
    for row_index, row in enumerate(rows, start=1):
        for column_index, column in enumerate(columns, start=1):
            position = f"{column}-{row}"
            if position not in values:
                continue
            table.cell(row_index, column_index).fill.solid()
            table.cell(row_index, column_index).fill.fore_color.rgb = _analysis_color(float(values[position]) * 100, thresholds)


def build_pptx_report(
    *,
    selected_sample: str,
    golden_sample: str,
    selected_version_label: str,
    golden_version_label: str,
    report_data: ReportData,
    ipf_images_by_direction: Optional[Dict[str, Dict[str, bytes]]] = None,
    ipf_images: Optional[Dict[str, bytes]] = None,
    ipf_legend: Optional[bytes] = None,
    scan_report_images: Optional[Dict[str, Dict[str, Any]]] = None,
    statistics_by_min_grain_size: Optional[Dict[int, Snapshot]] = None,
    analysis_data: Optional[Dict[str, Dict[str, float]]] = None,
    analysis_features: Optional[List[str]] = None,
    color_thresholds: Optional[List[float]] = None,
    overridden_positions: Optional[List[str]] = None,
    histogram_settings: Optional[Dict[str, Any]] = None,
) -> bytes:
    deck = TemplateDeck()
    selected = report_data.get(selected_sample, {})
    golden = report_data.get(golden_sample, {}) if golden_sample else {}
    directions = dict(ipf_images_by_direction or {})
    if not directions and ipf_images:
        directions["z"] = ipf_images
    overwritten = overridden_positions or []
    _add_scan_point_pages(
        deck,
        selected_sample,
        selected,
        scan_report_images or {},
        overwritten,
        ipf_legend,
        histogram_settings or {},
    )
    _add_ipf_pages(deck, selected_sample, selected_version_label, directions, ipf_legend)
    _add_grain_chart_pages(deck, selected, golden, selected_version_label, golden_version_label, histogram_settings or {})
    _add_complete_table(deck, selected_sample, selected_version_label, selected, overwritten)
    _add_orientation_tables(deck, selected, grains=False, overridden_positions=overwritten)
    _add_orientation_tables(deck, selected, grains=True, overridden_positions=overwritten)
    _add_statistics(deck, statistics_by_min_grain_size or {10: selected}, selected_sample, overwritten)
    _add_triangle_pages(deck, selected, golden)
    _add_orientation_lines(deck, selected, golden)
    _add_feature_analysis(deck, analysis_data or {}, selected_sample, analysis_features or [], color_thresholds or [30.0, 70.0])
    return deck.save()
