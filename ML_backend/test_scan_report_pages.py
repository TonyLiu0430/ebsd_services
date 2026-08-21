import io
import unittest

import numpy as np
from PIL import Image
from diffpy.structure import Lattice, Structure
from orix.crystal_map import Phase
from orix.quaternion import Orientation, symmetry
from orix.vector import Vector3d
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from scipy.ndimage import map_coordinates

from python_ebsd_figures import (
    RawRgbMap,
    TEXTURE_DENSITY_GRID_SIZE,
    _channel5_plot_coordinates,
    _grain_size_rgb,
    _inverse_pole_figure,
    _lambert_equal_area_density,
    _pole_figure,
    _raw_rgb_map,
)
from template_ppt_report import _deviation_legend, build_pptx_report


def _test_png() -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (160, 100), (120, 180, 210)).save(stream, format="PNG")
    return stream.getvalue()


def _scan_images() -> dict:
    image = _test_png()
    return {
        "band_contrast": image,
        "grain_map": image,
        "preferred_20": image,
        "preferred_15": image,
        "pole_figure": image,
        "inverse_pole_figure": image,
        "ipf_maps": {"X": image, "Y": image, "Z": image},
    }


class ScanPointPageTests(unittest.TestCase):
    def test_deviation_legend_uses_solid_orientation_colours(self) -> None:
        legend = Image.open(
            io.BytesIO(_deviation_legend(("[100]", "[110]", "[111]"), 20))
        ).convert("RGB")
        colours = legend.getcolors(maxcolors=legend.width * legend.height)
        self.assertIsNotNone(colours)
        colour_counts = {colour: count for count, colour in colours or []}

        self.assertGreater(colour_counts.get((242, 92, 92), 0), 100)
        self.assertGreater(colour_counts.get((89, 209, 99), 0), 100)
        self.assertGreater(colour_counts.get((102, 115, 230), 0), 100)

    def test_each_position_gets_two_pages_and_only_overwrite_is_reanalysis(self) -> None:
        item = {
            "grains": [4.0, 7.0, 11.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }
        report = build_pptx_report(
            selected_sample="Cu-U",
            golden_sample="",
            selected_version_label="2",
            golden_version_label="",
            report_data={"Cu-U": {"B-1": item, "A-2": item, "A-1": item}},
            scan_report_images={
                "B-1": _scan_images(),
                "A-2": _scan_images(),
                "A-1": _scan_images(),
            },
            overridden_positions=["A-1"],
        )
        presentation = Presentation(io.BytesIO(report))
        scan_titles = []
        for slide in presentation.slides:
            title = next((shape for shape in slide.shapes if shape.name == "EBSD_TITLE"), None)
            if title is not None and title.text.startswith("Cu-U "):
                scan_titles.append(title)

        self.assertEqual(
            [title.text for title in scan_titles],
            [
                "Cu-U A-1 (Reanalysis)",
                "Cu-U A-1 (Reanalysis)",
                "Cu-U A-2",
                "Cu-U A-2",
                "Cu-U B-1",
                "Cu-U B-1",
            ],
        )
        self.assertFalse(any("Reanalysis" in title.text for title in scan_titles if "A-1" not in title.text))

        reanalysis_runs = [
            run
            for title in scan_titles
            if "A-1" in title.text
            for paragraph in title.text_frame.paragraphs
            for run in paragraph.runs
            if run.text == " (Reanalysis)"
        ]
        self.assertEqual(len(reanalysis_runs), 2)
        self.assertTrue(all(run.font.color.rgb == RGBColor(0, 0, 255) for run in reanalysis_runs))

    def test_preferred_orientation_maps_end_above_ratio_tables(self) -> None:
        item = {
            "grains": [4.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }
        report = build_pptx_report(
            selected_sample="Cu-U",
            golden_sample="",
            selected_version_label="2",
            golden_version_label="",
            report_data={"Cu-U": {"A-1": item}},
            scan_report_images={"A-1": _scan_images()},
        )
        presentation = Presentation(io.BytesIO(report))
        texture_slide = next(
            slide
            for slide in presentation.slides
            if any(
                shape.has_text_frame and shape.text == "Preferred orientation ratio"
                for shape in slide.shapes
            )
        )
        preferred_maps = [
            shape
            for shape in texture_slide.shapes
            if shape.name in {"EBSD_PREFERRED_20", "EBSD_PREFERRED_15"}
        ]
        ratio_tables = [shape for shape in texture_slide.shapes if shape.has_table]

        self.assertEqual(len(preferred_maps), 2)
        self.assertEqual(len(ratio_tables), 2)
        table_top = min(shape.top for shape in ratio_tables)
        self.assertTrue(all(shape.top + shape.height < table_top for shape in preferred_maps))

    def test_texture_slide_has_deviation_legend_on_each_side(self) -> None:
        item = {
            "grains": [4.0],
            "orientation_labels": ["[100]", "[110]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }
        report = build_pptx_report(
            selected_sample="Cu-U",
            golden_sample="",
            selected_version_label="2",
            golden_version_label="",
            report_data={"Cu-U": {"A-1": item}},
            scan_report_images={"A-1": _scan_images()},
            ipf_legend=_test_png(),
        )
        presentation = Presentation(io.BytesIO(report))
        texture_slide = next(
            slide
            for slide in presentation.slides
            if any(
                shape.has_text_frame and shape.text == "Preferred orientation ratio"
                for shape in slide.shapes
            )
        )
        deviation_20 = next(shape for shape in texture_slide.shapes if shape.name == "EBSD_DEVIATION_20")
        deviation_15 = next(shape for shape in texture_slide.shapes if shape.name == "EBSD_DEVIATION_15")
        orientation_key = next(shape for shape in texture_slide.shapes if shape.name == "EBSD_ORIENTATION_KEY")

        self.assertEqual(deviation_20.shape_type, MSO_SHAPE_TYPE.PICTURE)
        self.assertEqual(deviation_15.shape_type, MSO_SHAPE_TYPE.PICTURE)
        self.assertLess(deviation_20.left + deviation_20.width, Inches(1.95))
        self.assertGreater(deviation_15.left, Inches(10.8))
        self.assertGreater(orientation_key.top, deviation_15.top + deviation_15.height)


class PythonTextureFigureTests(unittest.TestCase):
    def test_python_map_is_sent_to_cpp_as_raw_rgb(self) -> None:
        rgb = np.zeros((200, 300, 3), dtype=float)
        rgb[..., 1] = 1.0

        image = _raw_rgb_map(rgb, step_size=2.0, map_name="GS")

        self.assertIsInstance(image, RawRgbMap)
        self.assertEqual((image.width, image.height), (300, 200))
        self.assertEqual(image.x_step_um, 2.0)
        self.assertEqual(image.map_name, "GS")
        self.assertEqual(len(image.pixels), 300 * 200 * 3)
        self.assertEqual(image.pixels[:3], bytes((0, 255, 0)))

    def test_grain_map_colour_is_determined_by_equivalent_diameter(self) -> None:
        grains = np.array(
            [
                [1, 1, 0, 3, 3],
                [1, 1, 0, 3, 3],
                [0, 0, 0, 3, 3],
                [2, 2, 0, 3, 3],
                [2, 2, 0, 3, 3],
            ],
            dtype=int,
        )
        rgb = _grain_size_rgb(grains, step_size=2.0)

        # Grain IDs 1 and 2 have the same area and therefore the same colour.
        np.testing.assert_allclose(rgb[1, 1], rgb[4, 1])
        # The larger grain lies farther toward the red end of the size scale.
        self.assertGreater(rgb[3, 4, 0], rgb[1, 1, 0])
        self.assertLess(rgb[3, 4, 1], rgb[1, 1, 1])

    def test_channel5_pf_coordinate_basis(self) -> None:
        sample_directions = Vector3d(
            [
                [-1.0, 0.0, 0.0],  # Channel 5 X0, right
                [0.0, 0.0, -1.0],  # Channel 5 Z0, top
                [0.0, 1.0, 0.0],   # Projection axis, disk centre
            ]
        )
        np.testing.assert_allclose(
            _channel5_plot_coordinates(sample_directions),
            np.eye(3),
            atol=1e-12,
        )

    def test_spherical_density_has_no_azimuthal_seam_at_pole(self) -> None:
        projection_pole = Vector3d(np.tile([0.0, 1.0, 0.0], (500, 1)))
        density = _lambert_equal_area_density(projection_pole)
        angles = np.linspace(0, 2 * np.pi, 360, endpoint=False)
        center = (TEXTURE_DENSITY_GRID_SIZE - 1) / 2
        radius_pixels = 0.2 * TEXTURE_DENSITY_GRID_SIZE / 2
        ring = map_coordinates(
            density.data,
            [center + radius_pixels * np.sin(angles), center + radius_pixels * np.cos(angles)],
            order=1,
        )

        self.assertGreater(float(np.mean(ring)), 0)
        self.assertLess(float(np.std(ring) / np.mean(ring)), 0.02)

    def test_python_texture_pipeline_generates_pf_and_ipf_png(self) -> None:
        phase = Phase(
            name="Cu",
            point_group=symmetry.Oh,
            structure=Structure(
                lattice=Lattice(a=3.6, b=3.6, c=3.6, alpha=90, beta=90, gamma=90)
            ),
        )
        rng = np.random.default_rng(7)
        euler = rng.random((300, 3)) * np.array([2 * np.pi, np.pi, 2 * np.pi])
        orientations = Orientation.from_euler(
            euler,
            symmetry=symmetry.Oh,
            direction="lab2crystal",
        )

        pole_figure = _pole_figure(orientations, "cubic", phase)
        inverse_pole_figure = _inverse_pole_figure(orientations, symmetry.Oh)

        self.assertEqual(pole_figure[:8], b"\x89PNG\r\n\x1a\n")
        self.assertEqual(inverse_pole_figure[:8], b"\x89PNG\r\n\x1a\n")
        self.assertGreater(len(pole_figure), 10_000)
        self.assertGreater(len(inverse_pole_figure), 10_000)


if __name__ == "__main__":
    unittest.main()
