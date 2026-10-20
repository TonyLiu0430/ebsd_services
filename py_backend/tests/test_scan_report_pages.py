import io
import unittest
from unittest.mock import patch

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from app.analysis.figures import (
    RawRgbMap,
    _grain_size_rgb,
    _raw_rgb_map,
)
from app.reports.presentation import (
    _deviation_legend,
    _excel_column,
    _shared_charts,
    _shared_distribution_charts,
    _histogram_edges,
    _label_stride,
    _nice_percent_max,
    _sampled_label_indices,
    _triangle_annotation_mode,
    _triangle_percent_max,
    build_pptx_report,
)


def _test_png() -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (160, 100), (120, 180, 210)).save(stream, format="PNG")
    return stream.getvalue()


def _scan_images() -> dict:
    image = _test_png()
    return {
        "band_contrast": image,
        "grain_map": image,
        "grain_boundaries": image,
        "csl_map": image,
        "csl_summary": {"Σ3": 12.5},
        "kam_map": image,
        "kam_values": [0.5, 1.0, 1.5],
        "gos_map": image,
        "gos_values": [1.0, 2.0, 3.0],
        "preferred_20": image,
        "preferred_15": image,
        "pole_figure": image,
        "inverse_pole_figure": image,
        "ipf_maps": {"X": image, "Y": image, "Z": image},
    }


class ScanPointPageTests(unittest.TestCase):
    def test_distributions_are_rendered_by_shared_chart_code(self) -> None:
        charts = _shared_distribution_charts(
            [
                {
                    "sample": [0.35, 0.52, 0.71, 1.15, 1.42, 2.05],
                    "mode": "hist",
                    "histYMode": "count",
                    "fixedXMin": 0,
                    "autoBinCount": True,
                    "sampleColor": "#0ea5e9",
                },
                {
                    "sample": [0.41, 0.63, 0.88, 1.22, 1.71, 2.44],
                    "mode": "hist",
                    "histYMode": "count",
                    "fixedXMin": 0,
                    "autoBinCount": True,
                    "sampleColor": "#10b981",
                },
            ]
        )
        self.assertEqual(len(charts), 2)
        for chart in charts:
            source = Image.open(io.BytesIO(chart))
            self.assertEqual(source.format, "PNG")
            image = source.convert("RGBA")
            self.assertGreaterEqual(image.width, 1000)
            self.assertGreater(image.getchannel("A").getextrema()[1], 0)

    def test_deviation_legend_uses_solid_orientation_colours(self) -> None:
        legend = Image.open(
            io.BytesIO(_deviation_legend(("[100]", "[110]", "[111]"), 20))
        ).convert("RGB")
        colours = legend.getcolors(maxcolors=legend.width * legend.height)
        self.assertIsNotNone(colours)
        family_colours = [
            colour
            for _, colour in colours or []
            if max(colour) - min(colour) > 35 and colour != (0, 0, 0)
        ]
        self.assertGreater(len(family_colours), 80)

    def test_each_position_gets_two_default_pages_and_only_overwrite_is_reanalysis(self) -> None:
        item = {
            "grains": [4.0, 7.0, 11.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }
        report = build_pptx_report(
            selected_sample="Cu-U",
            golden_sample="",
            selected_version_label="2",
            golden_version_label="",
            report_data={"Cu-U": {"B-1": item, "A-2": item, "A-1": item}},
            scan_report_images={
                "B-1": _scan_images(),
                "A-2": _scan_images(),
                "A-1": _scan_images(),
            },
            overridden_positions=["A-1"],
        )
        presentation = Presentation(io.BytesIO(report))
        report_titles = {
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.name == "EBSD_TITLE"
        }
        self.assertIn("Orientation ratio", report_titles)
        self.assertNotIn("Orientation ratio by pixels", report_titles)
        self.assertNotIn("Orientation ratio by grains", report_titles)
        scan_titles = []
        for slide in presentation.slides:
            title = next((shape for shape in slide.shapes if shape.name == "EBSD_TITLE"), None)
            if title is not None and title.text.startswith("Cu-U "):
                scan_titles.append(title)

        self.assertEqual(
            [title.text for title in scan_titles],
            [
                "Cu-U A-1 (Reanalysis)",
                "Cu-U A-1 (Reanalysis)",
                "Cu-U A-2",
                "Cu-U A-2",
                "Cu-U B-1",
                "Cu-U B-1",
            ],
        )
        self.assertFalse(any("Reanalysis" in title.text for title in scan_titles if "A-1" not in title.text))

        reanalysis_runs = [
            run
            for title in scan_titles
            if "A-1" in title.text
            for paragraph in title.text_frame.paragraphs
            for run in paragraph.runs
            if run.text == " (Reanalysis)"
        ]
        self.assertEqual(len(reanalysis_runs), 2)
        self.assertTrue(all(run.font.color.rgb == RGBColor(0, 0, 255) for run in reanalysis_runs))

    def test_microstructure_page_only_exists_when_api_flag_is_true(self) -> None:
        item = {
            "grains": [4.0, 7.0, 11.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }

        def captions(report: bytes) -> set[str]:
            return {
                shape.text
                for slide in Presentation(io.BytesIO(report)).slides
                for shape in slide.shapes
                if shape.has_text_frame
            }

        with patch(
            "app.reports.presentation._shared_charts",
            side_effect=lambda charts, **_: [_test_png()] * len(charts),
        ), patch(
            "app.reports.presentation._triangle_chart", return_value=io.BytesIO(_test_png())
        ), patch(
            "app.reports.presentation._line_chart", return_value=io.BytesIO(_test_png())
        ):
            hidden = build_pptx_report(
                selected_sample="Cu-U",
                golden_sample="",
                selected_version_label="2",
                golden_version_label="",
                report_data={"Cu-U": {"A-1": item}},
                scan_report_images={"A-1": _scan_images()},
            )
            shown = build_pptx_report(
                selected_sample="Cu-U",
                golden_sample="",
                selected_version_label="2",
                golden_version_label="",
                report_data={"Cu-U": {"A-1": item}},
                scan_report_images={"A-1": _scan_images()},
                show_microstructure_page=True,
            )

        microstructure_captions = {
            "Grain boundaries",
            "Coincident site lattice (CSL)",
            "Kernel average misorientation (KAM)",
            "Grain orientation spread (GOS)",
            "KAM distribution",
            "GOS distribution",
        }
        self.assertTrue(microstructure_captions.isdisjoint(captions(hidden)))
        self.assertTrue(microstructure_captions.issubset(captions(shown)))

    def test_preferred_orientation_maps_end_above_ratio_tables(self) -> None:
        item = {
            "grains": [4.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }
        report = build_pptx_report(
            selected_sample="Cu-U",
            golden_sample="",
            selected_version_label="2",
            golden_version_label="",
            report_data={"Cu-U": {"A-1": item}},
            scan_report_images={"A-1": _scan_images()},
        )
        presentation = Presentation(io.BytesIO(report))
        texture_slide = next(
            slide
            for slide in presentation.slides
            if any(
                shape.has_text_frame and shape.text == "Preferred orientation ratio"
                for shape in slide.shapes
            )
        )
        preferred_maps = [
            shape
            for shape in texture_slide.shapes
            if shape.name in {"EBSD_PREFERRED_20", "EBSD_PREFERRED_15"}
        ]
        ratio_tables = [shape for shape in texture_slide.shapes if shape.has_table]

        self.assertEqual(len(preferred_maps), 2)
        self.assertEqual(len(ratio_tables), 2)
        table_top = min(shape.top for shape in ratio_tables)
        self.assertTrue(all(shape.top + shape.height < table_top for shape in preferred_maps))

    def test_texture_slide_has_deviation_legend_on_each_side(self) -> None:
        item = {
            "grains": [4.0],
            "orientation_labels": ["[100]", "[110]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }
        report = build_pptx_report(
            selected_sample="Cu-U",
            golden_sample="",
            selected_version_label="2",
            golden_version_label="",
            report_data={"Cu-U": {"A-1": item}},
            scan_report_images={"A-1": _scan_images()},
            ipf_legend=_test_png(),
        )
        presentation = Presentation(io.BytesIO(report))
        texture_slide = next(
            slide
            for slide in presentation.slides
            if any(
                shape.has_text_frame and shape.text == "Preferred orientation ratio"
                for shape in slide.shapes
            )
        )
        deviation_20 = next(shape for shape in texture_slide.shapes if shape.name == "EBSD_DEVIATION_20")
        deviation_15 = next(shape for shape in texture_slide.shapes if shape.name == "EBSD_DEVIATION_15")
        orientation_key = next(shape for shape in texture_slide.shapes if shape.name == "EBSD_ORIENTATION_KEY")

        self.assertEqual(deviation_20.shape_type, MSO_SHAPE_TYPE.PICTURE)
        self.assertEqual(deviation_15.shape_type, MSO_SHAPE_TYPE.PICTURE)
        self.assertLess(deviation_20.left + deviation_20.width, Inches(1.95))
        self.assertGreater(deviation_15.left, Inches(10.8))
        self.assertGreater(orientation_key.top, deviation_15.top + deviation_15.height)

    def test_texture_ratio_tables_use_compact_teal_header_style(self) -> None:
        item = {
            "grains": [4.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.122, 0.424, 0.25],
            "orientation_ratio(15%)": [0.0422, 0.481, 0.252],
        }
        report = build_pptx_report(
            selected_sample="Cu-U",
            golden_sample="",
            selected_version_label="2",
            golden_version_label="",
            report_data={"Cu-U": {"A-1": item}},
            scan_report_images={"A-1": _scan_images()},
        )
        presentation = Presentation(io.BytesIO(report))
        texture_slide = next(
            slide
            for slide in presentation.slides
            if any(
                shape.has_text_frame and shape.text == "Preferred orientation ratio"
                for shape in slide.shapes
            )
        )
        tables = [shape.table for shape in texture_slide.shapes if shape.has_table]

        self.assertEqual(len(tables), 2)
        for table in tables:
            self.assertEqual([table.cell(0, column).text for column in range(3)], ["<001>", "<011>", "<111>"])
            for column in range(3):
                header_run = table.cell(0, column).text_frame.paragraphs[0].runs[0]
                value_run = table.cell(1, column).text_frame.paragraphs[0].runs[0]
                self.assertTrue(header_run.font.bold)
                self.assertEqual(header_run.font.color.rgb, RGBColor(255, 255, 255))
                self.assertEqual(header_run.font.size.pt, 13)
                self.assertEqual(table.cell(0, column).fill.fore_color.rgb, RGBColor(70, 167, 193))
                self.assertEqual(value_run.font.size.pt, 12)
                self.assertEqual(table.cell(1, column).fill.fore_color.rgb, RGBColor(235, 247, 251))


class ReportAnnotationTests(unittest.TestCase):
    def test_histogram_x_max_rounds_page_data_up_to_nearest_ten(self) -> None:
        edges = _histogram_edges(
            [4.0, 17.5, 47.1],
            {"start_mode": "zero", "width_mode": "absolute", "width_absolute": 10.0},
        )
        self.assertEqual(edges[0], 0.0)
        self.assertEqual(edges[-1], 50.0)

    def test_axis_visibility_does_not_change_exported_chart_dimensions(self) -> None:
        common = {
            "sample": [5.0, 15.0, 25.0],
            "golden": [7.0, 17.0, 27.0],
            "mode": "hist",
            "histYMode": "ratio",
            "fixedXMin": 0,
            "fixedXMax": 30,
            "binWidth": 10,
            "canvasAspectRatio": 2.0,
            "fontScale": 1.7,
        }
        charts = _shared_charts(
            [
                {**common, "showXAxis": True, "showYAxis": True},
                {**common, "showXAxis": False, "showYAxis": False},
            ]
        )
        sizes = [Image.open(io.BytesIO(chart)).size for chart in charts]
        self.assertEqual(sizes[0], sizes[1])

    def test_shared_v8_rounds_the_whole_page_peak_up_to_tens(self) -> None:
        values = [5.0] * 43 + [15.0 + 10.0 * (index % 19) for index in range(57)]
        metadata = {}
        _shared_charts(
            [
                {
                    "sample": values,
                    "golden": [],
                    "mode": "hist",
                    "histYMode": "ratio",
                    "fixedXMin": 0,
                    "fixedXMax": 200,
                    "binWidth": 10,
                }
            ],
            shared_y_max=True,
            result_metadata=metadata,
        )
        self.assertEqual(metadata["sharedYMax"], 50)

    def test_histograms_show_bar_width_and_comparisons_identify_golden_data(self) -> None:
        selected_item = {
            "grains": [4.0, 7.0, 11.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.1, 0.2, 0.3],
            "orientation_ratio(15%)": [0.05, 0.1, 0.15],
        }
        golden_item = {
            **selected_item,
            "grains": [5.0, 8.0, 12.0],
        }
        report = build_pptx_report(
            selected_sample="Sample-Cu",
            golden_sample="Golden-Cu",
            selected_version_label="Sample-Cu-2",
            golden_version_label="Golden-Cu-3",
            report_data={
                "Sample-Cu": {"A-1": selected_item},
                "Golden-Cu": {"A-1": golden_item},
            },
            scan_report_images={"A-1": _scan_images()},
            histogram_settings={
                "start_mode": "min",
                "width_mode": "absolute",
                "width_absolute": 12.5,
            },
            analysis_data={"Sample-Cu": {"A-1": 0.25}},
            analysis_features=["grain size average"],
        )
        presentation = Presentation(io.BytesIO(report))

        grain_summary_slide = next(
            slide
            for slide in presentation.slides
            if any(
                shape.name == "EBSD_TITLE" and shape.text.startswith("Grain Size Summary")
                for shape in slide.shapes
            )
        )
        grain_summary_tables = [shape.table for shape in grain_summary_slide.shapes if shape.has_table]
        grain_summary_text = " ".join(
            cell.text
            for table in grain_summary_tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertEqual(len(grain_summary_tables), 3)
        self.assertIn("Grain Size Average (µm)", grain_summary_text)
        self.assertIn("Grain Size Max (µm)", grain_summary_text)
        self.assertIn("Grain Size Quantity", grain_summary_text)
        self.assertNotIn("Minimum", grain_summary_text)

        scan_histogram_note = next(
            shape
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.name == "EBSD_HISTOGRAM_NOTE"
        )
        self.assertEqual(scan_histogram_note.text, "Bar width: 12.5 µm")

        expected_comparison_titles = {
            "Grain histogram (Ratio)",
            "晶粒取向比例 - Misorientation 20%",
            "晶粒取向比例 - Misorientation 15%",
            "晶粒取向折線圖 - Misorientation 20%",
            "晶粒取向折線圖 - Misorientation 15%",
            "Feature analysis",
        }
        comparison_notes = {}
        for slide in presentation.slides:
            title = next((shape.text for shape in slide.shapes if shape.name == "EBSD_TITLE"), "")
            if title not in expected_comparison_titles:
                continue
            note = next(shape.text for shape in slide.shapes if shape.name == "EBSD_REPORT_NOTE")
            comparison_notes[title] = note

        self.assertEqual(set(comparison_notes), expected_comparison_titles)
        self.assertTrue(
            all("Golden data: Golden-Cu · Golden-Cu-3" in note for note in comparison_notes.values())
        )
        for title in expected_comparison_titles - {"Feature analysis"}:
            if "histogram" in title.lower():
                self.assertIn("Bar width: 12.5 µm", comparison_notes[title])


class ResponsiveMatrixLayoutTests(unittest.TestCase):
    @staticmethod
    def _snapshot(column_count: int, row_count: int) -> dict:
        item = {
            "grains": [5.0, 12.0, 24.0],
            "orientation_labels": ["[001]", "[011]", "[111]"],
            "orientation_ratio(20%)": [0.12, 0.27, 0.58],
            "orientation_ratio(15%)": [0.08, 0.19, 0.41],
        }
        return {
            f"{_excel_column(column + 1)}-{row + 1}": item
            for row in range(row_count)
            for column in range(column_count)
        }

    def test_arbitrary_n_by_m_matrices_stay_on_single_responsive_slides(self) -> None:
        image = _test_png()
        responsive_fonts = {}
        matrix_sizes = ((2, 7), (3, 3), (5, 4), (8, 9), (11, 6), (27, 1))
        for column_count, row_count in matrix_sizes:
            snapshot = self._snapshot(column_count, row_count)
            ipf_images = {position: image for position in snapshot}
            placeholder = io.BytesIO(image)
            with patch(
                "app.reports.presentation._shared_charts",
                side_effect=lambda charts, **_: [image] * len(charts),
            ), patch(
                "app.reports.presentation._triangle_chart", return_value=placeholder
            ), patch("app.reports.presentation._line_chart", return_value=placeholder):
                report = build_pptx_report(
                    selected_sample="Sample",
                    golden_sample="",
                    selected_version_label="Sample-1",
                    golden_version_label="",
                    report_data={"Sample": snapshot},
                    ipf_images_by_direction={"Z": ipf_images},
                )
            presentation = Presentation(io.BytesIO(report))
            titles = [
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.name == "EBSD_TITLE"
            ]
            self.assertEqual(sum("Orientation maps" in title for title in titles), 1)
            self.assertEqual(titles.count("Grain histogram (Ratio)"), 1)
            self.assertEqual(sum(title.startswith("晶粒取向比例") for title in titles), 2)
            orientation_slide = next(
                slide
                for slide in presentation.slides
                if any(shape.name == "EBSD_TITLE" and shape.text == "Orientation ratio" for shape in slide.shapes)
            )
            orientation_table = next(shape.table for shape in orientation_slide.shapes if shape.has_table)
            responsive_fonts[(column_count, row_count)] = (
                orientation_table.cell(2, 1).text_frame.paragraphs[0].runs[0].font.size.pt
            )
            complete_slide = next(
                slide
                for slide in presentation.slides
                if any(shape.name == "EBSD_TITLE" and shape.text == "Complete EBSD data" for shape in slide.shapes)
            )
            self.assertEqual(
                sum(shape.has_table for shape in complete_slide.shapes),
                2 if column_count * row_count > 40 else 1,
            )
            for slide in presentation.slides:
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, presentation.slide_width)
                    self.assertLessEqual(shape.top + shape.height, presentation.slide_height)
        self.assertGreater(responsive_fonts[(3, 3)], responsive_fonts[(5, 4)])
        self.assertGreater(responsive_fonts[(5, 4)], responsive_fonts[(8, 9)])

    def test_orientation_ratio_tables_do_not_show_ebsd_in_top_left_cell(self) -> None:
        image = _test_png()
        snapshot = self._snapshot(3, 3)
        placeholder = io.BytesIO(image)
        with patch(
            "app.reports.presentation._shared_charts",
            side_effect=lambda charts, **_: [image] * len(charts),
        ), patch(
            "app.reports.presentation._triangle_chart", return_value=placeholder
        ), patch("app.reports.presentation._line_chart", return_value=placeholder):
            report = build_pptx_report(
                selected_sample="Sample",
                golden_sample="",
                selected_version_label="Sample-1",
                golden_version_label="",
                report_data={"Sample": snapshot},
            )

        presentation = Presentation(io.BytesIO(report))
        orientation_slide = next(
            slide
            for slide in presentation.slides
            if any(
                shape.name == "EBSD_TITLE" and shape.text == "Orientation ratio"
                for shape in slide.shapes
            )
        )
        tables = [shape.table for shape in orientation_slide.shapes if shape.has_table]
        self.assertEqual(len(tables), 2)
        self.assertTrue(all(table.cell(0, 0).text == "" for table in tables))

    def test_xyz_ipf_directions_each_get_their_own_ppt_slide(self) -> None:
        image = _test_png()
        snapshot = self._snapshot(1, 1)
        placeholder = io.BytesIO(image)
        with patch(
            "app.reports.presentation._shared_charts",
            side_effect=lambda charts, **_: [image] * len(charts),
        ), patch(
            "app.reports.presentation._triangle_chart", return_value=placeholder
        ), patch("app.reports.presentation._line_chart", return_value=placeholder):
            report = build_pptx_report(
                selected_sample="Sample",
                golden_sample="",
                selected_version_label="Sample-1",
                golden_version_label="",
                report_data={"Sample": snapshot},
                ipf_images_by_direction={
                    "X": {"A-1": image},
                    "Y": {"A-1": image},
                    "Z": {"A-1": image},
                },
            )

        presentation = Presentation(io.BytesIO(report))
        orientation_map_titles = [
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.name == "EBSD_TITLE" and "Orientation maps" in shape.text
        ]
        self.assertEqual(
            orientation_map_titles,
            [
                "Sample · Orientation maps · IPF X",
                "Sample · Orientation maps · IPF Y",
                "Sample · Orientation maps · IPF Z",
            ],
        )

    def test_every_grain_chart_has_complete_axes_and_identical_canvas_settings(self) -> None:
        image = _test_png()
        snapshot = self._snapshot(3, 3)
        captured = []

        def render(charts, **_):
            captured.extend(charts)
            return [image] * len(charts)

        with patch("app.reports.presentation._shared_charts", side_effect=render), patch(
            "app.reports.presentation._triangle_chart", return_value=io.BytesIO(image)
        ), patch("app.reports.presentation._line_chart", return_value=io.BytesIO(image)):
            build_pptx_report(
                selected_sample="Sample",
                golden_sample="",
                selected_version_label="Sample-1",
                golden_version_label="",
                report_data={"Sample": snapshot},
            )

        self.assertEqual(len(captured), 9)
        self.assertTrue(all(chart["showXAxis"] for chart in captured))
        self.assertTrue(all(chart["showYAxis"] for chart in captured))
        self.assertEqual(len({chart["canvasAspectRatio"] for chart in captured}), 1)
        self.assertEqual(len({chart["fontScale"] for chart in captured}), 1)

    def test_dense_line_labels_are_sampled_from_arbitrary_column_count(self) -> None:
        self.assertEqual(_label_stride(11, 12), 1)
        self.assertEqual(_label_stride(27, 12), 3)
        self.assertEqual(_label_stride(53, 10), 6)
        self.assertEqual(_sampled_label_indices(27, 3), [0, 3, 6, 9, 12, 15, 18, 21, 26])

    def test_triangle_annotations_follow_cell_density_not_a_named_matrix_size(self) -> None:
        self.assertEqual(_triangle_annotation_mode(0.60, True), (True, True))
        self.assertEqual(_triangle_annotation_mode(0.30, True), (True, False))
        self.assertEqual(_triangle_annotation_mode(0.17, True), (False, False))
        self.assertEqual(_triangle_annotation_mode(0.60, False), (True, False))

    def test_triangle_scale_uses_fifty_floor_then_tens_without_headroom(self) -> None:
        self.assertEqual(_triangle_percent_max([]), 50.0)
        self.assertEqual(_triangle_percent_max([18.4, 49.9]), 50.0)
        self.assertEqual(_triangle_percent_max([50.0]), 50.0)
        self.assertEqual(_triangle_percent_max([50.1]), 60.0)
        self.assertEqual(_triangle_percent_max([60.0]), 60.0)
        self.assertEqual(_triangle_percent_max([60.1]), 70.0)

    def test_line_scale_rounds_page_maximum_to_tens_without_headroom(self) -> None:
        self.assertEqual(_nice_percent_max([]), 10.0)
        self.assertEqual(_nice_percent_max([9.9]), 10.0)
        self.assertEqual(_nice_percent_max([50.0]), 50.0)
        self.assertEqual(_nice_percent_max([50.1]), 60.0)
        self.assertEqual(_nice_percent_max([60.0]), 60.0)
        self.assertEqual(_nice_percent_max([60.1]), 70.0)


class PythonTextureFigureTests(unittest.TestCase):
    def test_python_map_is_sent_to_cpp_as_raw_rgb(self) -> None:
        rgb = np.zeros((200, 300, 3), dtype=float)
        rgb[..., 1] = 1.0

        image = _raw_rgb_map(rgb, step_size=2.0, map_name="GS")

        self.assertIsInstance(image, RawRgbMap)
        self.assertEqual((image.width, image.height), (300, 200))
        self.assertEqual(image.x_step_um, 2.0)
        self.assertEqual(image.map_name, "GS")
        self.assertEqual(len(image.pixels), 300 * 200 * 3)
        self.assertEqual(image.pixels[:3], bytes((0, 255, 0)))

    def test_grain_map_colour_is_determined_by_equivalent_diameter(self) -> None:
        grains = np.array(
            [
                [1, 1, 0, 3, 3],
                [1, 1, 0, 3, 3],
                [0, 0, 0, 3, 3],
                [2, 2, 0, 3, 3],
                [2, 2, 0, 3, 3],
            ],
            dtype=int,
        )
        rgb = _grain_size_rgb(grains, step_size=2.0)

        # Grain IDs 1 and 2 have the same area and therefore the same colour.
        np.testing.assert_allclose(rgb[1, 1], rgb[4, 1])
        # The larger grain lies farther toward the red end of the size scale.
        self.assertGreater(rgb[3, 4, 0], rgb[1, 1, 0])
        self.assertLess(rgb[3, 4, 1], rgb[1, 1, 1])

if __name__ == "__main__":
    unittest.main()
